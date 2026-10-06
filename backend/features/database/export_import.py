import os
import re

import webview

from .connect import mongo_db_name, _db_connections
from .crud import _SKIP, _build_mongo_filter, _build_search_clause, _build_where, _coerce_value
from .schema import fetch_schema


def _sql_literal(value, family):
    if value is None:
        return "NULL"
    if family in ("int", "float"):
        return str(value)
    if family == "bool":
        return "TRUE" if str(value).strip().lower() in ("true", "1", "t") else "FALSE"
    return "'" + str(value).replace("'", "''") + "'"


class DatabaseExportImportMixin:
    def db_export_table(self, path, table, filters=None, search=None):
        if not re.match(r"^[a-zA-Z0-9_$.-]+$", table or ""):
            return {"error": "Nom de table invalide"}
        conn, engine, error = self._db_connect(path)
        if error:
            return {"error": error}
        try:
            if engine == "mongo":
                db = conn[mongo_db_name(_db_connections[path]["uri"]) or conn.list_database_names()[0]]
                coll = db[table]
                mongo_filter = _build_mongo_filter(filters)
                docs = list(coll.find(mongo_filter))
                columns, seen = [], set()
                for doc in docs:
                    for key in doc.keys():
                        if key not in seen:
                            seen.add(key)
                            columns.append(key)
                rows = [[None if doc.get(c) is None else str(doc.get(c)) for c in columns] for doc in docs]
            else:
                schema_columns, _pk = fetch_schema(conn, engine, table)
                valid_columns = {c["name"] for c in schema_columns}
                quote = "`" if engine == "mysql" else '"'
                quoted = f"{quote}{table}{quote}"
                placeholder = "?" if engine == "sqlite" else "%s"
                where_sql, where_params = _build_where(engine, valid_columns, filters, placeholder)
                search_sql, search_params = _build_search_clause(engine, schema_columns, search, placeholder)
                if search_sql:
                    where_sql = f" WHERE {search_sql}" if not where_sql else f"{where_sql} AND{search_sql}"
                    where_params = where_params + search_params
                cur = conn.cursor()
                cur.execute(f"SELECT * FROM {quoted}{where_sql}", tuple(where_params))
                columns = [d[0] for d in cur.description]
                rows = [[None if v is None else str(v) for v in row] for row in cur.fetchall()]
                cur.close()
        except Exception as e:
            return {"error": str(e)}
        finally:
            conn.close()

        return self._write_xlsx_dialog(columns, rows, f"{table}.xlsx", table)

    def _write_xlsx_dialog(self, columns, rows, suggested_name, sheet_title):
        dest = webview.windows[0].create_file_dialog(
            webview.FileDialog.SAVE, save_filename=suggested_name, file_types=("Fichiers Excel (*.xlsx)",)
        )
        if not dest:
            return {"cancelled": True}
        dest = dest[0] if isinstance(dest, (list, tuple)) else dest
        if not dest.lower().endswith(".xlsx"):
            dest += ".xlsx"

        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = (sheet_title or "Sheet1")[:31]
        ws.append(columns)
        for row in rows:
            ws.append(row)
        try:
            wb.save(dest)
        except Exception as e:
            return {"error": str(e)}
        return {"path": dest, "count": len(rows)}

    def db_import_file(self, path, table):
        if not re.match(r"^[a-zA-Z0-9_$.-]+$", table or ""):
            return {"error": "Nom de table invalide"}
        # pywebview's filter parser (parse_file_type) only accepts
        # "[\w ]+(*.ext;*.ext2)" — a "/" or other punctuation in the label
        # makes it raise internally, which create_file_dialog swallows and
        # turns into a silent None (looked like the button did nothing at all).
        picked = webview.windows[0].create_file_dialog(
            webview.FileDialog.OPEN, file_types=("Fichiers CSV et Excel (*.csv;*.xlsx)", "Tous les fichiers (*.*)")
        )
        if not picked:
            return {"cancelled": True}
        src = picked[0] if isinstance(picked, (list, tuple)) else picked

        try:
            if src.lower().endswith(".xlsx"):
                from openpyxl import load_workbook

                wb = load_workbook(src, read_only=True, data_only=True)
                ws = wb.active
                rows_iter = ws.iter_rows(values_only=True)
                header = [str(c) if c is not None else "" for c in next(rows_iter, [])]
                file_rows = [[("" if c is None else c) for c in r] for r in rows_iter]
            else:
                import csv

                with open(src, "r", encoding="utf-8-sig", newline="") as f:
                    reader = csv.reader(f)
                    header = next(reader, [])
                    file_rows = list(reader)
        except Exception as e:
            return {"error": f"Lecture du fichier : {e}"}

        conn, engine, error = self._db_connect(path)
        if error:
            return {"error": error}
        try:
            if engine == "mongo":
                db = conn[mongo_db_name(_db_connections[path]["uri"]) or conn.list_database_names()[0]]
                docs = [
                    {h: v for h, v in zip(header, r) if h and v != ""}
                    for r in file_rows
                ]
                if not docs:
                    return {"error": "Aucune ligne à importer"}
                result = db[table].insert_many(docs)
                return {"inserted": len(result.inserted_ids), "skipped": 0, "errors": []}

            schema_columns, _pk_column = fetch_schema(conn, engine, table)
            by_name = {c["name"]: c for c in schema_columns}
            # Matched case-insensitively so a spreadsheet header like "Name"
            # still lines up with a lowercase "name" column.
            header_map = {}
            lower_by_name = {n.lower(): n for n in by_name}
            for h in header:
                match = lower_by_name.get((h or "").strip().lower())
                if match:
                    header_map[h] = match

            quote = "`" if engine == "mysql" else '"'
            quoted_table = f"{quote}{table}{quote}"
            placeholder = "?" if engine == "sqlite" else "%s"

            inserted, skipped, errors = 0, 0, []
            cur = conn.cursor()
            for i, raw_row in enumerate(file_rows):
                values = dict(zip(header, raw_row))
                cols, params, row_error = [], [], None
                for h, col_name in header_map.items():
                    col = by_name[col_name]
                    coerced, err = _coerce_value(values.get(h), col["family"], col["nullable"], col["has_default"])
                    if err:
                        row_error = f'ligne {i + 2}, "{col_name}" : {err}'
                        break
                    if coerced is _SKIP:
                        continue
                    cols.append(col_name)
                    params.append(coerced)
                if row_error:
                    skipped += 1
                    if len(errors) < 20:
                        errors.append(row_error)
                    continue
                if not cols:
                    skipped += 1
                    continue
                quoted_cols = ", ".join(f"{quote}{c}{quote}" for c in cols)
                placeholders = ", ".join([placeholder] * len(cols))
                try:
                    cur.execute(f"INSERT INTO {quoted_table} ({quoted_cols}) VALUES ({placeholders})", tuple(params))
                    inserted += 1
                except Exception as e:
                    skipped += 1
                    if len(errors) < 20:
                        errors.append(f"ligne {i + 2} : {e}")
            conn.commit()
            cur.close()
            return {"inserted": inserted, "skipped": skipped, "errors": errors}
        except Exception as e:
            conn.rollback()
            return {"error": str(e)}
        finally:
            conn.close()

    def db_export_schema_sql(self, path, include_data=False):
        conn, engine, error = self._db_connect(path)
        if error:
            return {"error": error}
        if engine == "mongo":
            return {"error": "Export SQL non disponible pour MongoDB."}
        try:
            tables_result = self.db_list_tables(path)
            if tables_result.get("error"):
                return tables_result
            quote = "`" if engine == "mysql" else '"'
            lines = [f"-- Export schéma{' + données' if include_data else ''} — {os.path.basename(path)}", ""]
            fk_statements = []
            for name in tables_result["tables"]:
                columns, pk_column = fetch_schema(conn, engine, name)
                col_defs = []
                for c in columns:
                    parts = [f"{quote}{c['name']}{quote}", c["type"].upper()]
                    if not c["nullable"]:
                        parts.append("NOT NULL")
                    col_defs.append(" ".join(parts))
                if pk_column:
                    col_defs.append(f"PRIMARY KEY ({quote}{pk_column}{quote})")
                lines.append(f"CREATE TABLE {quote}{name}{quote} (")
                lines.append("  " + ",\n  ".join(col_defs))
                lines.append(");")
                lines.append("")
                for c in columns:
                    if c.get("fk"):
                        fk_statements.append(
                            f"ALTER TABLE {quote}{name}{quote} ADD FOREIGN KEY ({quote}{c['name']}{quote}) "
                            f"REFERENCES {quote}{c['fk']['table']}{quote} ({quote}{c['fk']['column']}{quote});"
                        )
                if include_data:
                    cur = conn.cursor()
                    cur.execute(f"SELECT * FROM {quote}{name}{quote}")
                    col_names = [d[0] for d in cur.description]
                    quoted_cols = ", ".join(f"{quote}{n}{quote}" for n in col_names)
                    by_name = {c["name"]: c for c in columns}
                    for row in cur.fetchall():
                        vals = []
                        for col_name, v in zip(col_names, row):
                            vals.append(_sql_literal(v, by_name.get(col_name, {}).get("family", "text")))
                        lines.append(f"INSERT INTO {quote}{name}{quote} ({quoted_cols}) VALUES ({', '.join(vals)});")
                    cur.close()
                    lines.append("")
            if fk_statements:
                lines.append("-- Foreign keys")
                lines.extend(fk_statements)
                lines.append("")
        except Exception as e:
            return {"error": str(e)}
        finally:
            conn.close()

        dest = webview.windows[0].create_file_dialog(
            webview.FileDialog.SAVE, save_filename="schema.sql", file_types=("Fichiers SQL (*.sql)",)
        )
        if not dest:
            return {"cancelled": True}
        dest = dest[0] if isinstance(dest, (list, tuple)) else dest
        if not dest.lower().endswith(".sql"):
            dest += ".sql"
        try:
            with open(dest, "w", encoding="utf-8") as f:
                f.write("\n".join(lines))
        except Exception as e:
            return {"error": str(e)}
        return {"path": dest}
